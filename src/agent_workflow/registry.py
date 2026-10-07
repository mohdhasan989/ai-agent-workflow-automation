from pathlib import Path

from openpyxl import load_workbook

from .models import TestQuestion, WorkflowDefinition

DEFAULT_WORKBOOK = (
    Path(__file__).resolve().parents[2] / "data" / "workflows" / "AI_Agent_Workflow_Assessment.xlsx"
)

WORKFLOW_COLUMNS = [
    "Workflow_ID",
    "Workflow_Name",
    "Trigger",
    "Inputs",
    "Steps",
    "Decision_Logic",
    "Tools_Required",
    "Expected_Output",
]
TEST_QUESTION_COLUMNS = ["Workflow_ID", "Test_Request", "What_To_Check"]


class RegistryError(Exception):
    pass


def _text(value) -> str:
    return "" if value is None else str(value).strip()


def _split_list(value) -> list[str]:
    return [part.strip() for part in _text(value).split(";") if part.strip()]


def _split_steps(value) -> list[str]:
    return [step.strip() for step in _text(value).split("→") if step.strip()]


def _header_index(header_row, required_columns: list[str], sheet_name: str) -> dict[str, int]:
    index = {name: pos for pos, name in enumerate(header_row) if name is not None}
    missing = [column for column in required_columns if column not in index]
    if missing:
        raise RegistryError(f"Sheet '{sheet_name}' is missing columns: {', '.join(missing)}")
    return index


def _cell(row, index: dict[str, int], column: str):
    pos = index[column]
    return row[pos] if pos < len(row) else None


class WorkflowRegistry:
    def __init__(self, workbook_path: Path = DEFAULT_WORKBOOK):
        self.workbook_path = Path(workbook_path)
        self.workflows: dict[str, WorkflowDefinition] = {}
        self.test_questions: list[TestQuestion] = []

    def load(self) -> "WorkflowRegistry":
        if not self.workbook_path.exists():
            raise RegistryError(f"Workbook not found: {self.workbook_path}")
        workbook = load_workbook(self.workbook_path, data_only=True, read_only=True)
        try:
            self.workflows = self._load_workflows(workbook)
            self.test_questions = self._load_test_questions(workbook)
        finally:
            workbook.close()
        return self

    def get_all_workflows(self) -> list[WorkflowDefinition]:
        return list(self.workflows.values())

    def get_workflow(self, workflow_id: str) -> WorkflowDefinition:
        try:
            return self.workflows[workflow_id]
        except KeyError:
            raise RegistryError(f"Unknown Workflow_ID: {workflow_id}") from None

    def get_all_test_questions(self) -> list[TestQuestion]:
        return list(self.test_questions)

    def _load_workflows(self, workbook) -> dict[str, WorkflowDefinition]:
        if "Workflows" not in workbook.sheetnames:
            raise RegistryError("Sheet 'Workflows' not found in workbook")
        rows = workbook["Workflows"].iter_rows(values_only=True)
        try:
            header = next(rows)
        except StopIteration:
            raise RegistryError("Sheet 'Workflows' is empty") from None
        index = _header_index(header, WORKFLOW_COLUMNS, "Workflows")

        workflows: dict[str, WorkflowDefinition] = {}
        for row_number, row in enumerate(rows, start=2):
            if all(cell is None for cell in row):
                continue
            workflow_id = _text(_cell(row, index, "Workflow_ID"))
            workflow_name = _text(_cell(row, index, "Workflow_Name"))
            steps = _split_steps(_cell(row, index, "Steps"))
            if not workflow_id:
                raise RegistryError(f"Sheet 'Workflows' row {row_number}: Workflow_ID is empty")
            if not workflow_name:
                raise RegistryError(
                    f"Sheet 'Workflows' row {row_number}: Workflow_Name is empty ({workflow_id})"
                )
            if not steps:
                raise RegistryError(f"Sheet 'Workflows' row {row_number}: Steps is empty ({workflow_id})")
            if workflow_id in workflows:
                raise RegistryError(f"Duplicate Workflow_ID: {workflow_id} (row {row_number})")
            workflows[workflow_id] = WorkflowDefinition(
                workflow_id=workflow_id,
                workflow_name=workflow_name,
                trigger=_text(_cell(row, index, "Trigger")),
                inputs=_split_list(_cell(row, index, "Inputs")),
                steps=steps,
                decision_logic=_text(_cell(row, index, "Decision_Logic")),
                tools_required=_split_list(_cell(row, index, "Tools_Required")),
                expected_output=_text(_cell(row, index, "Expected_Output")),
            )
        return workflows

    def _load_test_questions(self, workbook) -> list[TestQuestion]:
        if "Test_Questions" not in workbook.sheetnames:
            raise RegistryError("Sheet 'Test_Questions' not found in workbook")
        rows = workbook["Test_Questions"].iter_rows(values_only=True)
        try:
            header = next(rows)
        except StopIteration:
            raise RegistryError("Sheet 'Test_Questions' is empty") from None
        index = _header_index(header, TEST_QUESTION_COLUMNS, "Test_Questions")

        questions: list[TestQuestion] = []
        for row_number, row in enumerate(rows, start=2):
            if all(cell is None for cell in row):
                continue
            workflow_id = _text(_cell(row, index, "Workflow_ID"))
            if not workflow_id:
                raise RegistryError(f"Sheet 'Test_Questions' row {row_number}: Workflow_ID is empty")
            questions.append(
                TestQuestion(
                    workflow_id=workflow_id,
                    test_request=_text(_cell(row, index, "Test_Request")),
                    what_to_check=_text(_cell(row, index, "What_To_Check")),
                )
            )
        return questions
