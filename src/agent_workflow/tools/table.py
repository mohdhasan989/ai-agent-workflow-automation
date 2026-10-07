import re
from pathlib import Path

from openpyxl import load_workbook

from .io import read_csv


def read_table(path, sheet=None, delimiter=","):
    path = Path(path)
    if path.suffix.lower() in (".xlsx", ".xlsm"):
        return _read_xlsx(path, sheet)
    return read_csv(str(path), delimiter=delimiter)


def _read_xlsx(path, sheet):
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        if sheet is None:
            sheet = workbook.sheetnames[0]
        rows = workbook[sheet].iter_rows(values_only=True)
        header = next(rows)
        columns = []
        for position, name in enumerate(header):
            if name is not None and str(name).strip():
                columns.append((position, str(name).strip()))
        result = []
        for row in rows:
            if all(value is None for value in row):
                continue
            result.append(
                {name: row[position] for position, name in columns if position < len(row)}
            )
        return result
    finally:
        workbook.close()


def detect_columns(rows):
    return list(rows[0].keys()) if rows else []


def normalize_columns(rows, rename=None):
    rename = rename or {}
    keys = list(rows[0].keys()) if rows else []
    mapping = {key: rename.get(key, _normalize_key(key)) for key in keys}
    return [{mapping.get(key, key): value for key, value in row.items()} for row in rows]


def _normalize_key(key):
    return re.sub(r"[-\s]+", "_", str(key).strip().lower())